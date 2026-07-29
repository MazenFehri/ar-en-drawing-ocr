import io
import itertools
import tempfile
import os
from typing import Iterator, Optional
from lxml import etree
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from models.elements import (
    Element, TextElement, SimpleShapeElement, PolylineShapeElement, ComplexShapeElement,
)
from utils.image_utils import ndarray_to_png_bytes

# A4 page geometry in EMU (914400 EMU = 1 inch), portrait orientation
_PAGE_W = 7_559_670
_PAGE_H = 10_692_720
_MARGIN = 914_400
_USE_W = _PAGE_W - 2 * _MARGIN
_USE_H = _PAGE_H - 2 * _MARGIN

# Minimum stroke width so a degenerate (zero-area) element is still visible
_MIN_VISIBLE_EMU = 91_440  # 0.1 inch
# ponytail: floor for a genuinely thin axis (a line) so we don't force it
# square; ceiling is "1 EMU is effectively invisible" — if renders show
# invisible hairlines, revisit with a real min-stroke-width heuristic.
_DEGENERATE_FLOOR_EMU = 1

# 12700 EMU = 1 point (the unit w:sz is expressed in, as half-points).
_EMU_PER_PT = 12_700
# A word's bbox height is its measured ink extent, not a typographic line
# box — there's no built-in ascender/descender headroom. 0.7 leaves ~30%
# slack so tall glyphs (Arabic ascenders, English capitals+descenders like
# "gh") don't kiss the box edges, while still using most of the box.
# ponytail: fixed ratio, not per-script metrics; if real fonts show
# clipping/looseness for a particular script, measure and adjust here.
_FONT_TO_BOX_HEIGHT_RATIO = 0.7
# ponytail: 4pt floor keeps degenerate/near-zero boxes legible-ish rather
# than invisible; 72pt ceiling stops a single huge detected box (e.g. a
# title) from producing absurd type. Both are round numbers, not measured —
# revisit if real drawings show either edge getting hit often.
_MIN_FONT_HALF_PT = 8    # 4pt
_MAX_FONT_HALF_PT = 144  # 72pt


def _font_sz_half_points(cy: int) -> int:
    """Font size (in half-points, what w:sz/w:szCs expect) for a box of
    height cy EMU. Keyed off height only — width doesn't matter (a long,
    short label must still get a short font, not a huge one)."""
    pt = (cy / _EMU_PER_PT) * _FONT_TO_BOX_HEIGHT_RATIO
    half_pt = round(pt * 2)
    return max(_MIN_FONT_HALF_PT, min(_MAX_FONT_HALF_PT, half_pt))

# DrawingML preset geometry names for each simple shape type
_WORD_SHAPE = {
    "circle": "ellipse",
    "ellipse": "ellipse",
    "triangle": "triangle",
    "rect": "rect",
    "square": "rect",
    "line": "line",
}

_HIGHLIGHT_MAP = {"yellow": "yellow", "red": "red"}

# Coordinate space a custGeom path is expressed in before DrawingML scales it to the
# shape extent. 100000 gives ~5 significant digits of path precision at any size.
_PATH_SPACE = 100_000

# Every anchor used to carry the same relativeHeight, which left stacking order
# up to document order — and reconstruct_layout emits text first, shapes second,
# so a shape whose bbox enclosed a label was painted *over* that label and hid
# it. Text inside a shape has to read on top of it, so give the graphics a lower
# band than the text and stop relying on list order for something this visible.
# The values are arbitrary but must keep their relative order; Word writes
# numbers in this range itself.
_Z_GRAPHIC = 251658240
_Z_TEXT = 251659264


def assemble_document(
    elements: list[Element],
    crop_images: Optional[dict[str, np.ndarray]] = None,
    page_aspect: Optional[float] = None,
    flow_text: bool = False,
) -> bytes:
    """Assemble a Word document with absolutely positioned elements.

    elements: list of TextElement, SimpleShapeElement, ComplexShapeElement
    crop_images: mapping of element id -> numpy BGR image (for ComplexShapeElement)
    page_aspect: source image width / height. When given, the drawing area is
        letterboxed to this aspect ratio (and the page flips to landscape if
        the source is wider than tall) so shapes and bboxes aren't stretched.
        When None, keeps the legacy full-page stretch mapping unchanged.
    flow_text: render ordinary flowing paragraphs instead of pinned textboxes.
        The floating layout is right for a drawing (a label belongs *at* the
        room it names) and wrong for a text-dense page — a printed Arabic
        worksheet comes out as scattered unreadable fragments. See
        _assemble_flowing. page_aspect is ignored in this mode: a flowing
        document has no source rect to letterbox into.
    Returns: .docx file as bytes
    """
    crop_images = crop_images or {}
    if flow_text:
        return _assemble_flowing(elements, crop_images)

    _counter = itertools.count(1)
    doc = Document()
    # Remove the default empty paragraph Word adds
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    geom = _page_geometry(doc, page_aspect)

    para = doc.add_paragraph()
    run = para.add_run()

    for el in elements:
        left, top, w, h = _emu_coords(el.bbox, geom)
        xml_str = None

        if isinstance(el, TextElement):
            xml_str = _text_box_xml(
                el.content, left, top, w, h,
                is_rtl=(el.language == "arabic"),
                highlight=el.highlight,
                _counter=_counter,
            )
        elif isinstance(el, SimpleShapeElement):
            prst = _WORD_SHAPE.get(el.shape, "rect")
            xml_str = _shape_xml(prst, left, top, w, h, _counter=_counter)
        elif isinstance(el, PolylineShapeElement):
            if len(el.points) >= 2:
                xml_str = _polyline_xml(el.points, left, top, w, h, _counter=_counter)
        elif isinstance(el, ComplexShapeElement):
            crop = crop_images.get(el.id)
            if crop is not None:
                xml_str = _image_xml(doc, ndarray_to_png_bytes(crop), left, top, w, h, _counter=_counter)

        if xml_str is not None:
            drawing_el = etree.fromstring(xml_str)
            run._r.append(drawing_el)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _assemble_flowing(elements: list[Element], crop_images: dict[str, np.ndarray]) -> bytes:
    """One ordinary Word paragraph per TextElement, in the order given.

    The elements arrive already in reading order — layout_reconstructor's
    _reading_order groups words into lines and reverses right-to-left ones —
    so this must NOT re-sort them.

    Nor must it touch the strings. utils/bidi.reshape_for_display exists for
    rendering to a raster; here the recogniser's output is already in Unicode
    logical order and Word runs its own bidi algorithm over a <w:bidi/>
    paragraph. Reshaping or reversing would be a second correction on top of a
    correct string, and the failure mode is silent: it still renders as Arabic,
    just backwards.
    """
    doc = Document()
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)

    for el in elements:
        if isinstance(el, TextElement):
            _add_flow_paragraph(doc, el.content, is_rtl=(el.language == "arabic"))
        elif isinstance(el, ComplexShapeElement):
            # A photo/chart inside a text document is real content, so it stays —
            # inline, at its position in the reading order, rather than pinned.
            crop = crop_images.get(el.id)
            if crop is not None:
                _add_flow_picture(doc, ndarray_to_png_bytes(crop))
        # SimpleShapeElement / PolylineShapeElement are dropped. On a text-dense
        # page these are table borders and answer-box rectangles, not diagram
        # content; floating them over flowing paragraphs recreates precisely the
        # scattered mess this mode exists to fix.
        # ponytail: ceiling is that no table structure survives — a bordered
        # table becomes a plain run of paragraphs, and a genuine diagram on a
        # mostly-text page loses its line art. Upgrade path is to cluster the
        # dropped rectangles into rows/columns and emit a real <w:tbl> instead
        # of discarding them, if a caller ever needs the grid back.

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_flow_paragraph(doc: Document, text: str, is_rtl: bool) -> None:
    """Append a normal <w:p>, RTL-marked when the text is Arabic.

    Built as raw XML and appended to an empty paragraph rather than driven
    through python-docx's API, because <w:bidi/> and <w:rtl/> have no accessor
    there and both are order-sensitive within their parent (w:bidi precedes
    w:jc in CT_PPr). Appending to a fresh empty w:p makes the order trivially
    correct; add_paragraph() also guarantees the w:p lands *before* the body's
    trailing w:sectPr, which a raw body.append() would not.
    """
    ns_w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    jc = "right" if is_rtl else "left"
    bidi_tag = "<w:bidi/>" if is_rtl else ""
    rtl_tag = "<w:rtl/>" if is_rtl else ""

    para = doc.add_paragraph()
    para._p.append(etree.fromstring(f'<w:pPr {ns_w}>{bidi_tag}<w:jc w:val="{jc}"/></w:pPr>'))
    para._p.append(etree.fromstring(
        f'<w:r {ns_w}><w:rPr>{rtl_tag}</w:rPr>'
        f'<w:t xml:space="preserve">{safe_text}</w:t></w:r>'
    ))


def _add_flow_picture(doc: Document, img_bytes: bytes) -> None:
    """Append the crop as an inline picture in its own paragraph."""
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(img_bytes)
        tmp_path = f.name
    try:
        shape = doc.add_picture(tmp_path)
    finally:
        os.unlink(tmp_path)
    # add_picture uses the image's native pixel size at its native DPI, so a
    # wide crop silently runs off the right margin. Scale down (never up) to the
    # usable text width, keeping the aspect ratio.
    if shape.width > _USE_W:
        shape.height = int(shape.height * _USE_W / shape.width)
        shape.width = _USE_W


def _page_geometry(doc: Document, page_aspect: Optional[float]) -> tuple[int, int, int, int]:
    """Return (box_x, box_y, box_w, box_h): the EMU rect elements map into.

    bbox coordinates are fractions of the source image, so the rect the
    fractions get multiplied into must itself have the source's aspect
    ratio, or shapes stretch (bug: circles rendered as ellipses).

    page_aspect is source width / height. None keeps the legacy behaviour
    (full A4-portrait usable area, no letterboxing) so old callers/tests
    are unaffected.
    """
    if page_aspect is None:
        return _MARGIN, _MARGIN, _USE_W, _USE_H

    page_w, page_h = _PAGE_W, _PAGE_H
    landscape = page_aspect > 1
    if landscape:
        page_w, page_h = _PAGE_H, _PAGE_W  # swap: landscape A4

    section = doc.sections[0]
    section.page_width = page_w
    section.page_height = page_h
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT

    use_w = page_w - 2 * _MARGIN
    use_h = page_h - 2 * _MARGIN

    # Letterbox: fit the source aspect ratio inside the usable area, centred.
    box_w = use_w
    box_h = int(box_w / page_aspect)
    if box_h > use_h:
        box_h = use_h
        box_w = int(box_h * page_aspect)

    box_x = _MARGIN + (use_w - box_w) // 2
    box_y = _MARGIN + (use_h - box_h) // 2
    return box_x, box_y, box_w, box_h


def _emu_coords(bbox, geom: tuple[int, int, int, int]) -> tuple[int, int, int, int]:
    box_x, box_y, box_w, box_h = geom
    left = box_x + int(bbox.x * box_w)
    top = box_y + int(bbox.y * box_h)
    width = int(bbox.w * box_w)
    height = int(bbox.h * box_h)
    if width < _DEGENERATE_FLOOR_EMU and height < _DEGENERATE_FLOOR_EMU:
        # Both axes collapsed (a "point") - inflate to a visible minimum.
        # A shape that's only thin on ONE axis (e.g. a horizontal line) must
        # keep that thin axis, or it renders as a fat box instead of a line.
        width = max(width, _MIN_VISIBLE_EMU)
        height = max(height, _MIN_VISIBLE_EMU)
    else:
        width = max(width, _DEGENERATE_FLOOR_EMU)
        height = max(height, _DEGENERATE_FLOOR_EMU)
    return left, top, width, height


def _next_id(_counter: Iterator[int]) -> int:
    return next(_counter)


def _anchor_wrap(left: int, top: int, cx: int, cy: int, inner_xml: str, _counter: Iterator[int],
                 z: int = _Z_TEXT) -> str:
    eid = _next_id(_counter)
    return (
        '<w:drawing xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" '
        f'relativeHeight="{z}" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1" '
        'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
        '<wp:simplePos x="0" y="0"/>'
        f'<wp:positionH relativeFrom="page"><wp:posOffset>{left}</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="page"><wp:posOffset>{top}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/>'
        '<wp:effectExtent l="0" t="0" r="0" b="0"/>'
        '<wp:wrapNone/>'
        f'<wp:docPr id="{eid}" name="Element{eid}"/>'
        '<wp:cNvGraphicFramePr/>'
        + inner_xml +
        '</wp:anchor></w:drawing>'
    )


def _text_box_xml(text: str, left: int, top: int, cx: int, cy: int,
                  is_rtl: bool, highlight: Optional[str], _counter: Iterator[int]) -> str:
    ns_a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ns_wps = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    ns_w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    bidi_tag = "<w:bidi/>" if is_rtl else ""
    hl_tag = (
        f'<w:highlight w:val="{_HIGHLIGHT_MAP[highlight]}"/>'
        if highlight and highlight in _HIGHLIGHT_MAP
        else ""
    )
    # Escape XML special chars in text
    safe_text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    sz = _font_sz_half_points(cy)
    return _anchor_wrap(left, top, cx, cy,
        f'<a:graphic {ns_a}>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp {ns_wps}>'
        '<wps:cNvSpPr txBox="1"/>'
        '<wps:spPr>'
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '<a:noFill/><a:ln><a:noFill/></a:ln>'
        '</wps:spPr>'
        '<wps:txbx>'
        f'<w:txbxContent {ns_w}>'
        f'<w:p><w:pPr>{bidi_tag}</w:pPr>'
        f'<w:r><w:rPr>{hl_tag}<w:sz w:val="{sz}"/><w:szCs w:val="{sz}"/></w:rPr>'
        f'<w:t xml:space="preserve">{safe_text}</w:t>'
        '</w:r></w:p>'
        '</w:txbxContent>'
        '</wps:txbx>'
        # wrap="none": a slightly-too-wide string overflows visibly rather
        # than wrapping to a 2nd line and getting clipped by box height —
        # for this pipeline, visible overflow beats silent data loss.
        # anchor="ctr": centres the single line vertically in the box so
        # it isn't glued to the top edge with the ~30% slack below it.
        # Insets zeroed: the box is the exact OCR ink extent already;
        # Word's default 0.1in/0.05in padding would eat most of a small box.
        # No autofit child: spAutoFit would resize the *shape* to fit the
        # text, defeating the positional fidelity this pipeline exists for;
        # normAutofit shrinks text by its own heuristic scale, fighting the
        # font size we just computed. Omitting is equivalent to noAutofit.
        '<wps:bodyPr wrap="none" anchor="ctr" lIns="0" tIns="0" rIns="0" bIns="0"/>'
        '</wps:wsp>'
        '</a:graphicData>'
        '</a:graphic>',
        _counter,
    )


def _shape_xml(prst: str, left: int, top: int, cx: int, cy: int, _counter: Iterator[int]) -> str:
    ns_a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ns_wps = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'
    return _anchor_wrap(left, top, cx, cy,
        f'<a:graphic {ns_a}>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp {ns_wps}>'
        '<wps:cNvSpPr/>'
        '<wps:spPr>'
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        f'<a:prstGeom prst="{prst}"><a:avLst/></a:prstGeom>'
        # Architectural line art: outline only, no default Word blue fill.
        # noFill/ln must come right after prstGeom or Word refuses the file.
        '<a:noFill/>'
        '<a:ln w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'
        '</wps:spPr>'
        '<wps:bodyPr/>'
        '</wps:wsp>'
        '</a:graphicData>'
        '</a:graphic>',
        _counter,
        z=_Z_GRAPHIC,
    )


def _polyline_xml(points: list[tuple[float, float]], left: int, top: int, cx: int, cy: int,
                  _counter: Iterator[int]) -> str:
    """An open path as a real DrawingML freeform (a:custGeom), not a picture.

    A connector/leader line has no preset geometry, and rasterising it was the bug:
    shapes belong in Word as shapes. Same outline-only styling and same z-band as
    _shape_xml, so it still sits below the text.
    """
    ns_a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ns_wps = 'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"'

    # The path declares its own coordinate space (a:path w/h) which DrawingML scales to
    # the shape's extent. A fixed space, rather than cx/cy, keeps the geometry sane when
    # an extent has been floored to a near-zero EMU by _emu_coords.
    def _pt(p: tuple[float, float]) -> str:
        x = int(round(min(max(p[0], 0.0), 1.0) * _PATH_SPACE))
        y = int(round(min(max(p[1], 0.0), 1.0) * _PATH_SPACE))
        return f'<a:pt x="{x}" y="{y}"/>'

    segments = f"<a:moveTo>{_pt(points[0])}</a:moveTo>" + "".join(
        f"<a:lnTo>{_pt(p)}</a:lnTo>" for p in points[1:]
    )
    return _anchor_wrap(left, top, cx, cy,
        f'<a:graphic {ns_a}>'
        '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
        f'<wps:wsp {ns_wps}>'
        '<wps:cNvSpPr/>'
        '<wps:spPr>'
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        # avLst/gdLst are empty but present: Word writes them and the schema orders
        # custGeom's children avLst, gdLst, ahLst, cxnLst, rect, pathLst.
        '<a:custGeom><a:avLst/><a:gdLst/>'
        f'<a:pathLst><a:path w="{_PATH_SPACE}" h="{_PATH_SPACE}" fill="none">'
        f'{segments}'
        '</a:path></a:pathLst></a:custGeom>'
        # No a:close: this is an open stroke, closing it would draw a phantom chord
        # back to the start point.
        '<a:noFill/>'
        '<a:ln w="12700"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln>'
        '</wps:spPr>'
        '<wps:bodyPr/>'
        '</wps:wsp>'
        '</a:graphicData>'
        '</a:graphic>',
        _counter,
        z=_Z_GRAPHIC,
    )


def _image_xml(doc: Document, img_bytes: bytes, left: int, top: int, cx: int, cy: int, _counter: Iterator[int]) -> str:
    """Add image to document part and return the drawing XML referencing it.

    Strategy:
    1. Write image bytes to a temp file.
    2. Use doc.add_picture() to register the image with the document part and get a relationship.
    3. Extract the rId from the inline drawing element that python-docx inserts.
    4. Remove the inline drawing element from the document (we only wanted the rId).
    5. Build and return a floating anchor XML that references the same rId.
    """
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as f:
        f.write(img_bytes)
        tmp_path = f.name
    try:
        # add_picture returns an InlineShape; it also appends a paragraph to the doc body
        inline_shape = doc.add_picture(tmp_path)
    finally:
        os.unlink(tmp_path)

    # python-docx appends the picture inside a new paragraph at the end of the body.
    # Extract the rId from the <a:blip r:embed="..."> element inside that drawing.
    NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    body = doc.element.body

    # The last paragraph added by add_picture
    last_para = body.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
    )[-1]

    # Find the blip element to get rId
    blip = last_para.find(f".//{{{NS_A}}}blip")
    rId = blip.get(f"{{{NS_R}}}embed") if blip is not None else None

    # If blip not found, fall back to getting the last relationship key
    if rId is None:
        rId = list(doc.part.rels.keys())[-1]

    # Remove the paragraph python-docx added (we only needed the rId)
    body.remove(last_para)

    ns_a = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    ns_pic = 'xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture"'
    ns_r = 'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
    eid = _next_id(_counter)
    return _anchor_wrap(left, top, cx, cy,
        f'<a:graphic {ns_a}>'
        '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
        f'<pic:pic {ns_pic}>'
        f'<pic:nvPicPr><pic:cNvPr id="{eid}" name="img{eid}"/><pic:cNvPicPr/></pic:nvPicPr>'
        f'<pic:blipFill><a:blip {ns_r} r:embed="{rId}"/>'
        '<a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
        '<pic:spPr>'
        f'<a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom>'
        '</pic:spPr>'
        '</pic:pic>'
        '</a:graphicData>'
        '</a:graphic>',
        _counter,
        z=_Z_GRAPHIC,
    )
